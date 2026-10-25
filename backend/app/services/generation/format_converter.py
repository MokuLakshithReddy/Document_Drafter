from pathlib import Path
import html
import docx
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib import colors
from app.core.errors import AppException, AppErrorCode
from app.core.logging import logger

class FormatConverter:
    @classmethod
    def convert_docx_to_pdf(cls, docx_path: Path, output_pdf_path: Path) -> Path:
        """
        Converts a DOCX file to a PDF file using ReportLab Platypus engine.
        Converts paragraphs, formatting runs, and tables.
        """
        if not docx_path.exists():
            raise AppException(code=AppErrorCode.DOCUMENT_NOT_FOUND, message=f"DOCX source file not found: {docx_path}")

        try:
            doc = docx.Document(docx_path)
            pdf = SimpleDocTemplate(
                str(output_pdf_path),
                pagesize=letter,
                rightMargin=54,
                leftMargin=54,
                topMargin=54,
                bottomMargin=54
            )

            styles = getSampleStyleSheet()
            normal_style = styles["Normal"]
            normal_style.fontSize = 11
            normal_style.leading = 15

            story = []

            for elem in doc.element.body:
                tag = elem.tag.split("}")[-1] if "}" in elem.tag else elem.tag
                if tag == "p":
                    # Find matching paragraph
                    for p in doc.paragraphs:
                        if p._p is elem:
                            para_html = cls._paragraph_to_html(p)
                            if para_html.strip():
                                story.append(Paragraph(para_html, normal_style))
                            else:
                                story.append(Spacer(1, 8))
                            story.append(Spacer(1, 4))
                            break

                elif tag == "tbl":
                    # Find matching table
                    for t in doc.tables:
                        if t._tbl is elem:
                            table_flowable = cls._table_to_flowable(t, normal_style)
                            if table_flowable:
                                story.append(table_flowable)
                                story.append(Spacer(1, 10))
                            break

            if not story:
                story.append(Paragraph("Empty Document", normal_style))

            pdf.build(story)
            return output_pdf_path

        except Exception as e:
            logger.error(f"DOCX to PDF conversion failed: {e}")
            raise AppException(
                code=AppErrorCode.GENERATION_ERROR,
                message=f"DOCX to PDF conversion failed: {str(e)}"
            )

    @classmethod
    def _paragraph_to_html(cls, paragraph) -> str:
        parts = []
        for r in paragraph.runs:
            text = html.escape(r.text)
            if not text:
                continue
            if r.bold:
                text = f"<b>{text}</b>"
            if r.italic:
                text = f"<i>{text}</i>"
            if r.underline:
                text = f"<u>{text}</u>"
            parts.append(text)
        return "".join(parts)

    @classmethod
    def _table_to_flowable(cls, table, style) -> Table:
        table_data = []
        for row in table.rows:
            row_data = []
            for cell in row.cells:
                cell_text = cell.text.strip()
                escaped = html.escape(cell_text).replace("\n", "<br/>")
                row_data.append(Paragraph(escaped, style))
            table_data.append(row_data)

        if not table_data:
            return None

        t = Table(table_data)
        t.setStyle(TableStyle([
            ('BACKGROUND', (0,0), (-1,0), colors.whitesmoke),
            ('ALIGN', (0,0), (-1,-1), 'LEFT'),
            ('FONTNAME', (0,0), (-1,-1), 'Helvetica'),
            ('FONTSIZE', (0,0), (-1,-1), 10),
            ('BOTTOMPADDING', (0,0), (-1,-1), 6),
            ('TOPPADDING', (0,0), (-1,-1), 6),
            ('GRID', (0,0), (-1,-1), 0.5, colors.grey),
        ]))
        return t
