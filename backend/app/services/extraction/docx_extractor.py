from pathlib import Path
from typing import List
import docx
from app.services.extraction.models import ExtractedDocument, ExtractedUnit, ExtractedRun
from app.core.logging import logger

class DocxExtractor:
    @staticmethod
    def extract(file_path: Path, document_id: str) -> ExtractedDocument:
        doc = docx.Document(file_path)
        units: List[ExtractedUnit] = []
        raw_lines: List[str] = []

        total_paragraphs = 0
        total_tables = len(doc.tables)
        has_headers_footers = False

        # 1. Process Headers from Sections
        for s_idx, section in enumerate(doc.sections):
            try:
                header = section.header
                if header and not header.is_linked_to_previous:
                    for p_idx, p in enumerate(header.paragraphs):
                        text = p.text
                        if text and text.strip():
                            has_headers_footers = True
                            runs = [
                                ExtractedRun(
                                    run_index=r_idx,
                                    text=r.text,
                                    bold=bool(r.bold),
                                    italic=bool(r.italic),
                                    underline=bool(r.underline),
                                    font_name=r.font.name if r.font else None,
                                    font_size_pt=float(r.font.size.pt) if r.font and r.font.size else None
                                )
                                for r_idx, r in enumerate(p.runs)
                            ]
                            unit = ExtractedUnit(
                                container_id=f"header[{s_idx}]/p[{p_idx}]",
                                kind="header",
                                full_text=text,
                                runs=runs
                            )
                            units.append(unit)
                            raw_lines.append(text)
            except Exception as e:
                logger.debug(f"Header extraction note: {e}")

        # 2. Process Body Paragraphs
        for p_idx, p in enumerate(doc.paragraphs):
            text = p.text
            total_paragraphs += 1
            if text and text.strip():
                runs = [
                    ExtractedRun(
                        run_index=r_idx,
                        text=r.text,
                        bold=bool(r.bold),
                        italic=bool(r.italic),
                        underline=bool(r.underline),
                        font_name=r.font.name if r.font else None,
                        font_size_pt=float(r.font.size.pt) if r.font and r.font.size else None
                    )
                    for r_idx, r in enumerate(p.runs)
                ]
                unit = ExtractedUnit(
                    container_id=f"body/p[{p_idx}]",
                    kind="paragraph",
                    full_text=text,
                    runs=runs
                )
                units.append(unit)
                raw_lines.append(text)

        # 3. Process Tables
        for t_idx, table in enumerate(doc.tables):
            for r_idx, row in enumerate(table.rows):
                for c_idx, cell in enumerate(row.cells):
                    for p_idx, p in enumerate(cell.paragraphs):
                        text = p.text
                        if text and text.strip():
                            runs = [
                                ExtractedRun(
                                    run_index=run_idx,
                                    text=r.text,
                                    bold=bool(r.bold),
                                    italic=bool(r.italic),
                                    underline=bool(r.underline),
                                    font_name=r.font.name if r.font else None,
                                    font_size_pt=float(r.font.size.pt) if r.font and r.font.size else None
                                )
                                for run_idx, r in enumerate(p.runs)
                            ]
                            unit = ExtractedUnit(
                                container_id=f"table[{t_idx}]/r[{r_idx}]/c[{c_idx}]/p[{p_idx}]",
                                kind="table_cell",
                                full_text=text,
                                runs=runs
                            )
                            units.append(unit)
                            raw_lines.append(text)

        # 4. Process Footers
        for s_idx, section in enumerate(doc.sections):
            try:
                footer = section.footer
                if footer and not footer.is_linked_to_previous:
                    for p_idx, p in enumerate(footer.paragraphs):
                        text = p.text
                        if text and text.strip():
                            has_headers_footers = True
                            runs = [
                                ExtractedRun(
                                    run_index=r_idx,
                                    text=r.text,
                                    bold=bool(r.bold),
                                    italic=bool(r.italic),
                                    underline=bool(r.underline),
                                    font_name=r.font.name if r.font else None,
                                    font_size_pt=float(r.font.size.pt) if r.font and r.font.size else None
                                )
                                for r_idx, r in enumerate(p.runs)
                            ]
                            unit = ExtractedUnit(
                                container_id=f"footer[{s_idx}]/p[{p_idx}]",
                                kind="footer",
                                full_text=text,
                                runs=runs
                            )
                            units.append(unit)
                            raw_lines.append(text)
            except Exception as e:
                logger.debug(f"Footer extraction note: {e}")

        raw_text = "\n".join(raw_lines)
        return ExtractedDocument(
            document_id=document_id,
            format="docx",
            units=units,
            raw_text=raw_text,
            total_paragraphs=total_paragraphs,
            total_tables=total_tables,
            has_headers_footers=has_headers_footers,
            is_scanned=False
        )
