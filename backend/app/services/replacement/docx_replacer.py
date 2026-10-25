from pathlib import Path
from typing import List, Dict, Any, Tuple
import docx
from app.core.errors import AppException, AppErrorCode
from app.core.logging import logger

class DocxReplacer:
    @classmethod
    def apply_replacements(
        cls,
        working_file_path: Path,
        replacements: List[Dict[str, Any]]
    ) -> Dict[str, Any]:
        """
        Applies approved replacements to the working copy DOCX file at the specified container locators.
        Replacements format:
        [
            {
                "occurrence_id": str,
                "container_id": str, # e.g. "body/p[2]", "table[0]/r[1]/c[0]/p[0]", "header[0]/p[0]"
                "original_text": str,
                "new_text": str,
                "char_start": Optional[int],
                "char_end": Optional[int],
            }
        ]
        """
        # 1. Overlap detection check
        cls._check_overlapping_edits(replacements)

        doc = docx.Document(working_file_path)
        applied_count = 0
        skipped_count = 0
        details = []

        # Group replacements by container_id
        replacements_by_container: Dict[str, List[Dict[str, Any]]] = {}
        for r in replacements:
            cid = r.get("container_id", "")
            if cid:
                replacements_by_container.setdefault(cid, []).append(r)

        # 1. Body Paragraphs
        for p_idx, p in enumerate(doc.paragraphs):
            cid = f"body/p[{p_idx}]"
            if cid in replacements_by_container:
                reps = replacements_by_container[cid]
                app, skp, det = cls._apply_to_paragraph(p, reps, cid)
                applied_count += app
                skipped_count += skp
                details.extend(det)

        # 2. Tables
        for t_idx, table in enumerate(doc.tables):
            for r_idx, row in enumerate(table.rows):
                for c_idx, cell in enumerate(row.cells):
                    for p_idx, p in enumerate(cell.paragraphs):
                        cid = f"table[{t_idx}]/r[{r_idx}]/c[{c_idx}]/p[{p_idx}]"
                        if cid in replacements_by_container:
                            reps = replacements_by_container[cid]
                            app, skp, det = cls._apply_to_paragraph(p, reps, cid)
                            applied_count += app
                            skipped_count += skp
                            details.extend(det)

        # 3. Headers & Footers
        for s_idx, section in enumerate(doc.sections):
            try:
                header = section.header
                if header and not header.is_linked_to_previous:
                    for p_idx, p in enumerate(header.paragraphs):
                        cid = f"header[{s_idx}]/p[{p_idx}]"
                        if cid in replacements_by_container:
                            reps = replacements_by_container[cid]
                            app, skp, det = cls._apply_to_paragraph(p, reps, cid)
                            applied_count += app
                            skipped_count += skp
                            details.extend(det)
            except Exception as e:
                logger.debug(f"Header replacement error: {e}")

            try:
                footer = section.footer
                if footer and not footer.is_linked_to_previous:
                    for p_idx, p in enumerate(footer.paragraphs):
                        cid = f"footer[{s_idx}]/p[{p_idx}]"
                        if cid in replacements_by_container:
                            reps = replacements_by_container[cid]
                            app, skp, det = cls._apply_to_paragraph(p, reps, cid)
                            applied_count += app
                            skipped_count += skp
                            details.extend(det)
            except Exception as e:
                logger.debug(f"Footer replacement error: {e}")

        # Check for any replacements whose container was never found
        handled_occ_ids = {d["occurrence_id"] for d in details}
        for r in replacements:
            if r.get("occurrence_id") not in handled_occ_ids:
                skipped_count += 1
                details.append({
                    "occurrence_id": r.get("occurrence_id"),
                    "status": "missed",
                    "container": r.get("container_id", "unknown"),
                    "reason": "Container locator not found in document"
                })

        # Save changes to the working file
        doc.save(str(working_file_path))

        return {
            "applied_count": applied_count,
            "skipped_count": skipped_count,
            "details": details
        }

    @classmethod
    def _check_overlapping_edits(cls, replacements: List[Dict[str, Any]]):
        """
        Rejects conflicting or overlapping edits targeting the same container.
        """
        by_container: Dict[str, List[Dict[str, Any]]] = {}
        for r in replacements:
            cid = r.get("container_id", "")
            if not cid:
                continue
            by_container.setdefault(cid, []).append(r)

        for cid, reps in by_container.items():
            if len(reps) <= 1:
                continue
            for i in range(len(reps)):
                for j in range(i + 1, len(reps)):
                    r1 = reps[i]
                    r2 = reps[j]
                    s1 = r1.get("char_start")
                    e1 = r1.get("char_end")
                    s2 = r2.get("char_start")
                    e2 = r2.get("char_end")
                    if s1 is not None and e1 is not None and s2 is not None and e2 is not None:
                        if max(s1, s2) < min(e1, e2):
                            raise AppException(
                                code=AppErrorCode.OVERLAPPING_EDIT,
                                message=f"Conflicting or overlapping edits detected in container '{cid}' between '{r1.get('original_text')}' and '{r2.get('original_text')}'",
                                details={"container_id": cid, "edit_1": r1.get("occurrence_id"), "edit_2": r2.get("occurrence_id")}
                            )

    @classmethod
    def _apply_to_paragraph(
        cls,
        paragraph,
        reps: List[Dict[str, Any]],
        cid: str
    ) -> Tuple[int, int, List[Dict[str, Any]]]:
        """
        Applies a list of replacements to a paragraph.
        Replacements are sorted in descending order of char_start so earlier offsets are never invalidated.
        Handles phrases within a single run and phrases split across multiple formatting runs.
        """
        applied = 0
        skipped = 0
        details = []

        # Sort descending by char_start so right-to-left replacements do not displace earlier offsets
        sorted_reps = sorted(
            reps,
            key=lambda x: x.get("char_start") if x.get("char_start") is not None else -1,
            reverse=True
        )

        for rep in sorted_reps:
            orig = rep.get("original_text", "")
            new_val = rep.get("new_text", "")
            c_start = rep.get("char_start")
            c_end = rep.get("char_end")
            occ_id = rep.get("occurrence_id", "")

            if not orig:
                continue

            success = cls._replace_exact_span(paragraph, orig, new_val, c_start, c_end)
            if success:
                applied += 1
                details.append({"occurrence_id": occ_id, "status": "applied", "container": cid})
            else:
                skipped += 1
                details.append({"occurrence_id": occ_id, "status": "missed", "container": cid})

        return applied, skipped, details

    @classmethod
    def _replace_exact_span(
        cls,
        paragraph,
        original_text: str,
        new_text: str,
        char_start: int = None,
        char_end: int = None
    ) -> bool:
        runs = paragraph.runs
        if not runs:
            return False

        # Build current run spans
        run_spans = []
        cur_offset = 0
        for r in runs:
            length = len(r.text)
            run_spans.append((cur_offset, cur_offset + length, r))
            cur_offset += length

        full_para_text = "".join(r.text for r in runs)

        # Determine target match interval
        match_start = None
        match_end = None

        if char_start is not None and char_end is not None:
            # Validate if original_text is at [char_start, char_end]
            if char_start < len(full_para_text) and full_para_text[char_start:char_end] == original_text:
                match_start = char_start
                match_end = char_end
            else:
                # Check for small local offset drift (+/- 4 chars)
                search_min = max(0, char_start - 4)
                search_max = min(len(full_para_text), char_end + 4)
                local_sub = full_para_text[search_min:search_max]
                local_idx = local_sub.find(original_text)
                if local_idx != -1:
                    match_start = search_min + local_idx
                    match_end = match_start + len(original_text)
        else:
            # Only search globally if locator offsets were not provided
            idx = full_para_text.find(original_text)
            if idx != -1:
                match_start = idx
                match_end = idx + len(original_text)

        if match_start is None:
            return False

        # Find overlapping runs
        overlapping_runs = []
        for r_start, r_end, r_obj in run_spans:
            if r_end > match_start and r_start < match_end:
                overlapping_runs.append((r_start, r_end, r_obj))

        if not overlapping_runs:
            return False

        first_start, first_end, first_run = overlapping_runs[0]
        prefix_len = max(0, match_start - first_start)
        first_prefix = first_run.text[:prefix_len]

        last_start, last_end, last_run = overlapping_runs[-1]
        suffix_start_in_last = max(0, match_end - last_start)
        last_suffix = last_run.text[suffix_start_in_last:] if suffix_start_in_last < len(last_run.text) else ""

        if len(overlapping_runs) == 1:
            first_run.text = first_prefix + new_text + last_suffix
        else:
            first_run.text = first_prefix + new_text
            for _, _, mid_run in overlapping_runs[1:-1]:
                mid_run.text = ""
            last_run.text = last_suffix

        return True
