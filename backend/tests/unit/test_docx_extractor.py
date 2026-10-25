import pytest
from pathlib import Path
from tests.fixtures.fixture_builder import create_complex_docx
from app.services.extraction.docx_extractor import DocxExtractor
from app.services.replacement.docx_replacer import DocxReplacer
import docx

def test_docx_extraction_and_locators(tmp_path: Path):
    doc_path = tmp_path / "complex_test.docx"
    create_complex_docx(doc_path)

    extracted = DocxExtractor.extract(doc_path, "doc_123")
    assert extracted.format == "docx"
    assert extracted.total_tables == 1
    assert extracted.has_headers_footers is True
    assert len(extracted.units) > 0

    # Verify container IDs exist
    container_ids = [u.container_id for u in extracted.units]
    assert any("header" in cid for cid in container_ids)
    assert any("footer" in cid for cid in container_ids)
    assert any("table" in cid for cid in container_ids)
    assert any("body" in cid for cid in container_ids)

def test_docx_multi_run_spanning_replacement(tmp_path: Path):
    doc_path = tmp_path / "multi_run.docx"
    create_complex_docx(doc_path)

    # In create_complex_docx:
    # "Alex " is run 1 (bold), "Morgan" is run 2 (bold)
    replacements = [
        {
            "occurrence_id": "occ_001",
            "container_id": "body/p[0]",
            "original_text": "Alex Morgan",
            "new_text": "Samantha Vance"
        }
    ]

    res = DocxReplacer.apply_replacements(doc_path, replacements)
    assert res["applied_count"] == 1

    # Verify replacement in document
    modified_doc = docx.Document(doc_path)
    body_text = "\n".join(p.text for p in modified_doc.paragraphs)
    assert "Samantha Vance" in body_text
    assert "Alex Morgan" not in body_text

def test_docx_table_replacement(tmp_path: Path):
    doc_path = tmp_path / "table_rep.docx"
    create_complex_docx(doc_path)

    replacements = [
        {
            "occurrence_id": "occ_table_01",
            "container_id": "table[0]/r[2]/c[1]/p[0]",
            "original_text": "$92,000",
            "new_text": "$115,000"
        }
    ]

    res = DocxReplacer.apply_replacements(doc_path, replacements)
    assert res["applied_count"] == 1

    modified_doc = docx.Document(doc_path)
    table_cell_text = modified_doc.tables[0].rows[2].cells[1].paragraphs[0].text
    assert table_cell_text == "$115,000"

def test_docx_unicode_and_punctuation(tmp_path: Path):
    doc_path = tmp_path / "unicode_test.docx"
    create_complex_docx(doc_path)

    unicode_val = "O'Connor & Sons — «€150,000» <Special & Safe>"
    replacements = [
        {
            "occurrence_id": "occ_uni",
            "container_id": "table[0]/r[2]/c[1]/p[0]",
            "original_text": "$92,000",
            "new_text": unicode_val
        }
    ]

    res = DocxReplacer.apply_replacements(doc_path, replacements)
    assert res["applied_count"] == 1

    modified_doc = docx.Document(doc_path)
    table_cell_text = modified_doc.tables[0].rows[2].cells[1].paragraphs[0].text
    assert table_cell_text == unicode_val
